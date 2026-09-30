"""Fills the official CollegeName_TeamName_Submission template -> docs/SRMIST_QueueQueueing.pptx.

The template (12 sections) stays intact: its masters, title slide design and section titles are kept; only
content is added. The template file itself is not committed (it lives in .tmp/template/template.pptx).
Run from the repo root:  .venv/Scripts/python docs/deck_src/build_template_deck.py
Member emails / video link come from docs/deck_src/team.json.
"""
import json
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / ".tmp" / "template" / "template.pptx"
OUT = ROOT / "docs" / "SRMIST_QueueQueueing.pptx"
ICONS = ROOT / ".tmp" / "icons"
TEAM = json.loads((Path(__file__).parent / "team.json").read_text(encoding="utf-8"))
TAG = "PRISM_GENAI_HACKATHON_Y2026"

# template palette (purple titles/accent, dark ink, muted grey) + blue/teal from its network art
PURPLE, VIOLET, INK, MUTED = "704EA6", "6D28D9", "14142B", "63637E"
SOFT, SOFT2, BLUE, TEAL = "F1ECFB", "F7F5FC", "2F5BD8", "2A9D8F"
GREEN, GREENSOFT, AMBER, AMBERSOFT, RED, REDSOFT, LINE = "1E9E5A", "E3F4EA", "C77700", "FDF1DE", "B3261E", "FBE9E7", "E2DCF0"
PRESENTER = {1: 0, 2: 0, 3: 0, 4: 1, 5: 1, 6: 1, 7: 2, 8: 2, 9: 2, 10: 3, 11: 3, 12: 3}


def rgb(h): return RGBColor.from_string(h)


def drop_body_placeholder(slide):
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx != 0:          # keep the title, drop the empty body box
            ph._element.getparent().remove(ph._element)


def box(slide, x, y, w, h, fill=SOFT2, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def text(slide, x, y, w, h, paras, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         font="Calibri", bullets=False, space_after=4, italic=False):
    """paras: str or list of str / (str, {size,color,bold,italic}) / list-of-runs [(str, opts), ...]."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [para]
        for r in runs:
            t, o = (r, {}) if isinstance(r, str) else r
            run = p.add_run()
            run.text = t
            f = run.font
            f.name = o.get("font", font); f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic); f.color.rgb = rgb(o.get("color", color))
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.2)))); pPr.set("indent", str(-Emu(Inches(0.2))))
            for tagname in ("a:buNone", "a:buChar", "a:buAutoNum"):
                for old in pPr.findall(qn(tagname)):
                    pPr.remove(old)
            bu = etree.SubElement(pPr, qn("a:buChar")); bu.set("char", "•")
    return tb


def icon(slide, key, x, y, d=0.6, fill=VIOLET):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = rgb(fill); c.line.fill.background(); c.shadow.inherit = False
    pad = d * 0.23
    slide.shapes.add_picture(str(ICONS / f"{key}_w.png"), Inches(x + pad), Inches(y + pad), Inches(d - 2 * pad), Inches(d - 2 * pad))


def arrow(slide, x, y, w=0.22, h=0.26, color=VIOLET):
    a = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    a.rotation = 90
    a.fill.solid(); a.fill.fore_color.rgb = rgb(color); a.line.fill.background(); a.shadow.inherit = False


def notes(slide, body):
    who = TEAM["members"][PRESENTER[slide_no(slide)]]["name"]
    slide.notes_slide.notes_text_frame.text = f"Presenter: {who}\n{body}"


_SLIDES = []


def slide_no(slide):
    return _SLIDES.index(slide) + 1


def table(slide, x, y, w, col_w, rows, header_fill=VIOLET, size=12, row_h=0.42):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    t = shp.table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08); cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(header_fill if i == 0 else ("FFFFFF" if i % 2 else SOFT2))
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.text = ""
            r = p.add_run(); r.text = str(val)
            r.font.name = "Calibri"; r.font.size = Pt(size); r.font.bold = i == 0 or (j >= 2 and i > 0 and len(row) == 4)
            r.font.color.rgb = rgb("FFFFFF" if i == 0 else INK)
    return t


def main() -> None:
    prs = Presentation(str(TEMPLATE))
    _SLIDES.extend(prs.slides)
    S = list(prs.slides)
    m = TEAM["members"]
    repo_short = TEAM["repo"].replace("https://", "")
    video = TEAM.get("video") or "link in README (YouTube, unlisted)"

    # ---------- 1. Title: fill the template's own fields ----------
    s = S[0]
    info = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Theme ID"))
    info.width = Inches(6.75); info.height = Inches(3.3)
    member = lambda i: f"{m[i]['name']} (Reg. No. {m[i]['reg']})" + (f", {m[i]['email']}" if m[i]["email"] else "") + ("  [primary member]" if m[i].get("primary") else "")
    values = ["Theme ID - 02 · Smart Guided Troubleshooting Engine · Project: FixFlow",
              f"Team Name - {TEAM['team']}", f"College Name - {TEAM['college']}",
              f"Member Name & Email 1- {member(0)}", f"Member Name & Email 2- {member(1)}",
              f"Member Name & Email 3- {member(2)}", f"Member Name & Email 4- {member(3)}",
              f"Submission Github link - {TEAM['repo']}"]
    for p, v in zip(info.text_frame.paragraphs, values):
        runs = p.runs
        runs[0].text = v
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        runs[0].font.size = Pt(13.5)
    notes(s, "Introduce the team and Theme 02. One line: FixFlow turns a vague device complaint into a grounded, one-tap, verifiable fix.")

    # ---------- 2. Theme ----------
    s = S[1]; drop_body_placeholder(s)
    text(s, 0.92, 1.72, 11.5, 0.4, [[("Theme 02 · Smart Guided Troubleshooting Engine", {"bold": True, "color": VIOLET, "size": 18})]])
    box(s, 0.92, 2.25, 6.55, 4.65, SOFT)
    icon(s, "quote", 1.15, 2.45, 0.6)
    text(s, 1.95, 2.48, 5.3, 0.9, "“My screen flickers and the battery dies fast.”", size=19, bold=True, italic=True, color=INK)
    text(s, 1.2, 3.45, 6.05, 3.4, [
        "Customers describe symptoms, not Settings screens.",
        "A support agent reads long knowledge articles, picks the steps that apply and orders them safely, about 15 minutes per scenario.",
        "The user then hunts through nested Settings menus to act on the advice.",
        "Our task: turn the complaint (+ the knowledge article) into clean, ordered steps as JSON, each Settings step with an exact one-tap deeplink, served in under 300 ms for known issues.",
    ], size=14, bullets=True, space_after=7)
    for i, (big, small, ic) in enumerate([("≈15 min", "manual triage per scenario", "clock"), ("Millions", "of support interactions", "headset"), ("10k+", "scenarios in production", "layer")]):
        y = 2.25 + i * 1.6
        box(s, 7.8, y, 4.62, 1.42, "FFFFFF", LINE)
        icon(s, ic, 8.05, y + 0.38, 0.66, PURPLE)
        text(s, 8.95, y + 0.2, 3.3, 0.6, big, size=30, bold=True, color=INK)
        text(s, 8.95, y + 0.82, 3.3, 0.4, small, size=13, color=MUTED)
    notes(s, "Problem in our own words: vague complaints, 15-minute manual triage, users lost in Settings. The deliverable is a REST API returning structured JSON with deeplinks.")

    # ---------- 3. Existing solutions & gaps ----------
    s = S[2]; drop_body_placeholder(s)
    rows = [["Approach today", "How it works", "Gap"],
            ["Manual agent triage", "Agent reads knowledge articles and writes steps by hand", "≈15 min per case; inconsistent order; no deeplinks"],
            ["FAQ / keyword search", "Returns whole articles that match words", "User still has to extract steps and find the screens"],
            ["Generic LLM chatbot", "Free-form answer from model memory", "Invents steps and web links; wrong or parent-menu screens; not auditable"],
            ["Rule / decision trees", "Hand-built flows per issue", "Brittle to new wording; does not scale to 10k+ scenarios"],
            ["Static deeplink lists", "Links per topic", "No link between a complaint and the exact screen or toggle"]]
    table(s, 0.92, 1.8, 11.5, [2.6, 4.2, 4.7], rows, size=12.5, row_h=0.5)
    box(s, 0.92, 5.05, 11.5, 1.8, SOFT)
    icon(s, "star", 1.15, 5.3, 0.62, VIOLET)
    text(s, 2.0, 5.22, 10.2, 0.4, "What is missing, and what FixFlow is designed for", size=16, bold=True, color=VIOLET)
    text(s, 2.0, 5.65, 10.2, 1.2, [
        "Grounded: only steps that exist in the reference article  ·  Exact: the right screen and the right on/off operation from the official catalog",
        "Safe: least disruptive first, restart/reset last  ·  Instant: repeats and paraphrases answered from a validated cache  ·  Verifiable: the fix can be checked",
    ], size=13.5, space_after=6)
    notes(s, "Walk through the table: every existing approach fails on at least one of grounding, exact screen, safe order, speed or scale.")

    # ---------- 4. Solution & architecture ----------
    s = S[3]; drop_body_placeholder(s)
    text(s, 1.0, 1.62, 11.4, 0.45, [[("FixFlow: ", {"bold": True, "color": VIOLET}), ("a REST engine that returns a schema-valid, ordered, deeplinked plan as JSON. The model proposes; code decides.", {})]], size=15)
    stages = [("route", "REST API", "POST /v1/troubleshoot\nGET /health", "[4]"), ("db", "Fast-path cache", "L1 exact + semantic\nL2 normalised complaint", "[3]"),
              ("search", "Query enrichment", "normalise, split intents,\n8–10 variations", "[0]"), ("brain", "Grounded extraction", "LLM cites numbered\nsentences; code verifies", "[1]"),
              ("link", "Deeplink mapping", "Feature Graph, hybrid\nBM25 + embeddings", "[2]"), ("shield", "Contract gate", "schema + every rule\nchecked in code", "✓")]
    for i, (ic, h, d, tag) in enumerate(stages):
        x = 0.92 + i * 1.95
        fill = AMBERSOFT if i == 1 else SOFT
        box(s, x, 2.25, 1.72, 2.25, fill)
        icon(s, ic, x + 0.56, 2.38, 0.6, AMBER if i == 1 else VIOLET)
        text(s, x + 0.08, 2.3, 0.5, 0.3, tag, size=10, bold=True, color=MUTED)
        text(s, x + 0.06, 3.05, 1.6, 0.5, h, size=13.5, bold=True, align=PP_ALIGN.CENTER)
        text(s, x + 0.06, 3.55, 1.6, 0.9, d, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 5:
            arrow(s, x + 1.72, 3.2, 0.2, 0.24)
    lanes = [("brain", "LLM layer", "Gemini Flash-Lite, auto-selected. Model fallback chain, retry with backoff, 8 s online budget, JSON-schema output, temperature 0."),
             ("cog", "Offline rules path", "No key or provider outage → rules parser with the same mapper and gate. Never fails; always JSON."),
             ("db", "Data", "Official kit: 578-entry deeplink catalog, 20 reference articles. Pre-validated plan cache (SQLite) ships in the Docker image.")]
    for i, (ic, h, d) in enumerate(lanes):
        x = 0.92 + i * 3.9
        box(s, x, 4.8, 3.7, 2.05, "FFFFFF", LINE)
        icon(s, ic, x + 0.2, 4.98, 0.5, PURPLE)
        text(s, x + 0.85, 5.03, 2.7, 0.4, h, size=15, bold=True)
        text(s, x + 0.2, 5.6, 3.35, 1.2, d, size=11.5, color=MUTED)
    notes(s, "Architecture left to right. Stress the two guarantees: steps must cite the article, deeplinks are copied verbatim from the catalog. Mention the offline fallback.")

    # ---------- 5. Demo & walkthrough ----------
    s = S[4]; drop_body_placeholder(s)
    s.shapes.add_picture(str(ROOT / "docs" / "img" / "ui_demo.png"), Inches(0.92), Inches(1.75), Inches(6.9), Inches(4.83))
    text(s, 0.92, 6.62, 6.9, 0.3, "Live UI served by the API: phone view (left) and engine view (right)", size=10.5, italic=True, color=MUTED)
    steps = [("mic", "Describe the problem", "Type, speak, or pick one of the 20 official complaints."),
             ("list", "Get an ordered plan", "Settings toggles → checks → service → restart/reset last."),
             ("check", "Tap → Verify", "Open applies the setting; the validation deeplink confirms it (“Touch sensitivity is True”)."),
             ("route", "Guided mode", "One step at a time: “Fixed?” or “Next”."),
             ("chart", "Engine view", "Cache / LLM / rules path, latency, cost, dropped unsupported steps.")]
    for i, (ic, h, d) in enumerate(steps):
        y = 1.78 + i * 0.93
        icon(s, ic, 8.12, y, 0.5, GREEN if ic == "check" else VIOLET)
        text(s, 8.8, y - 0.03, 3.65, 0.35, h, size=14, bold=True)
        text(s, 8.8, y + 0.32, 3.65, 0.55, d, size=11.5, color=MUTED)
    text(s, 8.12, 6.45, 4.3, 0.5, [[("Run: ", {"bold": True, "color": VIOLET}), ("docker compose up → localhost:8000", {})], [("Video: ", {"bold": True, "color": VIOLET}), (video, {})]], size=11)
    notes(s, "Show the demo flow: sample complaint, ordered plan, Open then Verify, guided mode, and the engine view proving the cache hit and $0 cost.")

    # ---------- 6. Tools & tech stack ----------
    s = S[5]; drop_body_placeholder(s)
    stack = [("python", "Python 3.12", "FastAPI · Uvicorn · Pydantic v2 (official schema.py, unchanged)"),
             ("brain", "Gemini API", "Flash-Lite (3.x) via REST, JSON-schema output, provider-agnostic adapter"),
             ("search", "Retrieval", "fastembed BAAI/bge-small-en-v1.5 (ONNX, CPU) · rank-bm25 · NumPy"),
             ("db", "Storage", "SQLite plan cache with embedded keys · JSON feature graph"),
             ("docker", "Delivery", "Docker + docker-compose · model and cache baked in · /health"),
             ("flask", "Quality", "pytest (47 hermetic tests) · evaluation harness · ablation · gold labels"),
             ("mobile", "Demo UI", "Static HTML/JS served at / · Web Speech API for voice input"),
             ("rocket", "Built with", "Claude Code as the engineering assistant")]
    for i, (ic, h, d) in enumerate(stack):
        x = 0.92 + (i % 2) * 5.85; y = 1.8 + (i // 2) * 1.27
        box(s, x, y, 5.6, 1.08, "FFFFFF", LINE)
        icon(s, ic, x + 0.22, y + 0.23, 0.62, VIOLET)
        text(s, x + 1.05, y + 0.16, 4.4, 0.35, h, size=15, bold=True)
        text(s, x + 1.05, y + 0.52, 4.4, 0.5, d, size=12, color=MUTED)
    notes(s, "CPU-only stack; runs with or without an API key. Point out the official schema is imported unchanged and 47 tests guard every contract rule.")

    # ---------- 7. Impact & use case ----------
    s = S[6]; drop_body_placeholder(s)
    uses = [("mobile", "Customer self-service", "Inside a voice assistant or support app: the user says the problem and gets one-tap fixes with verification."),
            ("headset", "Support-agent assist", "The agent pastes the complaint; FixFlow returns an ordered, cited plan in milliseconds instead of ≈15 minutes."),
            ("layer", "Knowledge base at scale", "Compile every knowledge article once into validated plans; serve millions of repeats from the cache at $0.")]
    for i, (ic, h, d) in enumerate(uses):
        x = 0.92 + i * 3.9
        box(s, x, 1.8, 3.7, 2.55, SOFT)
        icon(s, ic, x + 0.25, 2.0, 0.66, VIOLET)
        text(s, x + 0.25, 2.8, 3.25, 0.4, h, size=16, bold=True)
        text(s, x + 0.25, 3.22, 3.25, 1.1, d, size=12.5, color=MUTED)
    stats = [("≈15 min → ms", "manual triage vs cached plan"), ("100%", "catalog-valid deeplinks (no wrong screen)"), ("$0", "per repeat / paraphrase (cache)"), ("0", "invented web links (scrubbed + verified)")]
    for i, (big, small) in enumerate(stats):
        x = 0.92 + i * 2.93
        text(s, x, 4.75, 2.75, 0.7, big, size=28, bold=True, color=VIOLET if i != 2 else GREEN)
        text(s, x, 5.45, 2.75, 0.6, small, size=12.5, color=MUTED)
    text(s, 0.92, 6.25, 11.5, 0.6, "Who benefits: customers (faster fixes, fewer service visits), support teams (less handling time, consistent answers), the product team (one reusable mapping for 10k+ scenarios).", size=13, italic=True, color=INK)
    notes(s, "Three use cases and the impact numbers. Emphasise reusability: compile once, serve millions.")

    # ---------- 8. Innovation, results & limitations ----------
    s = S[7]; drop_body_placeholder(s)
    cols = [("Innovation", VIOLET, SOFT, ["Citation-verified extraction: code rejects steps not in the article",
                                            "Action- and depth-aware deeplinks (exact screen, right operation)",
                                            "Closed-loop Tap → Verify with validation deeplinks",
                                            "Trust-tiered, symptom-guarded two-level cache",
                                            "Outage-proof: model fallback chain + offline rules"]),
            ("Results", GREEN, GREENSOFT, ["Gates 100%: schema, rules, 0 URL leaks, catalog-valid links (D1 + D2)",
                                          "Deeplink precision 100%; relevance 2.00 (D1) / 1.71 (D2) of 2",
                                          "Step-accuracy proxy 2.20 (D1) / 2.88 (D2) of 3; abstention 90% / 100%",
                                          "Cache P95: 0 ms exact, 48 ms unseen paraphrase; 81.8% hit",
                                          "Ablation: hybrid 1.76 rel / 100% precision vs rules 1.29 / 58%"]),
            ("Limitations", RED, REDSOFT, ["Provider outage: the LLM built 12 of 20 official plans; the rest used rules",
                                             "Paraphrases: every hit had the right symptom; 59% the exact same article",
                                             "Offline relevance is coarse; LLM sometimes drops a section",
                                             "Gold labels by one annotator; D2 used once for error analysis",
                                             "Contract puts service escalation before restarts"])]
    for i, (h, col, fill, items) in enumerate(cols):
        x = 0.92 + i * 3.9
        box(s, x, 2.2, 3.7, 4.7, fill)                 # title wraps to two lines on this slide
        text(s, x + 0.25, 2.36, 3.2, 0.4, h, size=18, bold=True, color=col)
        text(s, x + 0.25, 2.88, 3.25, 3.95, items, size=12.5, bullets=True, space_after=8)
    notes(s, "Be honest here: results are strong on the automated gates and deeplinks; limitations include the provider outage and the paraphrase article ambiguity. Full tables in metrics.md.")

    # ---------- 9. What's next ----------
    s = S[8]; drop_body_placeholder(s)
    road = [("layer", "Compile the full base", "Run the batch compiler over all 10k+ articles into the pre-validated cache, with a human review tier."),
            ("chart", "Learn the best order", "Telemetry-driven ordering: which step fixes each symptom most often."),
            ("check", "Real device verification", "Execute Tap → Verify against live device state through the validation deeplinks."),
            ("globe", "Multilingual + on-device", "Hindi / Hinglish intake via LLM enrichment; small on-device model for the offline path.")]
    for i, (ic, h, d) in enumerate(road):
        x = 0.92 + i * 2.93
        box(s, x, 2.0, 2.7, 3.6, SOFT)
        text(s, x + 0.2, 2.12, 1.0, 0.35, f"Step {i + 1}", size=11, bold=True, color=MUTED)
        icon(s, ic, x + 1.0, 2.5, 0.7, VIOLET)
        text(s, x + 0.2, 3.4, 2.3, 0.45, h, size=15, bold=True, align=PP_ALIGN.CENTER)
        text(s, x + 0.2, 3.9, 2.3, 1.6, d, size=12, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 3:
            arrow(s, x + 2.72, 3.6, 0.18, 0.22)
    text(s, 0.92, 5.95, 11.5, 0.8, "Worklet fit: the engine, evaluation harness and batch compiler already exist; the next step is scale and live-device integration, not a rewrite.", size=14, italic=True, color=INK)
    notes(s, "Roadmap from prototype to worklet, four steps.")

    # ---------- 10. Brownie points (differentiation) ----------
    s = S[9]; drop_body_placeholder(s)
    diff = [("flask", "Measured, not claimed", "Held-out paraphrases, 4-domain scenarios, gold labels, and a scripted 3-way ablation, with honest limits."),
            ("sync", "Outage-proof, proven", "During a real Gemini 503 outage, 30 of 32 cold requests fell back to rules and every one stayed contract-valid."),
            ("check", "Closes the loop", "Validation deeplinks confirm the fix actually happened, not just a link to a screen."),
            ("shield", "Contract-exact", "Official schema unchanged, every rule enforced in code, brand-neutral outputs, 47 hermetic tests."),
            ("docker", "Zero-friction to judge", "docker compose up: model and validated cache baked in; runs with no API key.")]
    for i, (ic, h, d) in enumerate(diff):
        y = 1.8 + i * 1.02
        icon(s, ic, 0.95, y + 0.08, 0.56, VIOLET)
        text(s, 1.7, y + 0.02, 5.4, 0.35, h, size=14.5, bold=True)
        text(s, 1.7, y + 0.38, 5.4, 0.6, d, size=11.5, color=MUTED)
    cd = CategoryChartData()
    cd.categories = ["Deeplink relevance (0–2)", "Link precision (0–1)"]
    cd.add_series("Hybrid BM25 + dense (ours)", (1.76, 1.00))
    cd.add_series("Pure keyword rules", (1.29, 0.58))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.45), Inches(1.75), Inches(4.95), Inches(4.4), cd)
    ch = gf.chart
    ch.has_title = True; ch.chart_title.text_frame.text = "Ablation: same extraction, different mapper"
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
    ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(11)
    for ser, col in zip(ch.plots[0].series, (VIOLET, "B9A7E8")):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(col)
    pl = ch.plots[0]; pl.has_data_labels = True
    pl.data_labels.number_format = "0.00"; pl.data_labels.number_format_is_linked = False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END; pl.data_labels.font.size = Pt(11)
    va = ch.value_axis; va.minimum_scale = 0; va.maximum_scale = 2.2; va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("E8E4F2"); va.tick_labels.font.size = Pt(10)
    ch.category_axis.tick_labels.font.size = Pt(11)
    text(s, 7.45, 6.25, 4.95, 0.6, "Full-LLM mapping baseline implemented (eval/ablation.py); not completed due to the provider outage.", size=10, italic=True, color=MUTED)
    notes(s, "Differentiation: measured evaluation, proven fallback, closed-loop verification, contract exactness, one-command run. Point at the ablation chart.")

    # ---------- 11. Checklist ----------
    s = S[10]
    body = next(sh for sh in s.placeholders if sh.placeholder_format.idx != 0)
    answers = [f"Working prototype code — public GitHub repo: Y  ({repo_short}, tag {TAG})",
               "README with reproducible setup instructions: Y  (docker compose up, or Python 3.12 + 3 commands)",
               f"Demo video, max 5 minutes: {video}",
               "Presentation file (PPT or PDF): Y  (docs/SRMIST_QueueQueueing.pptx and .pdf)"]
    for p, v in zip(body.text_frame.paragraphs, answers):
        runs = p.runs
        runs[0].text = v
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        for r in p.runs:
            r.font.size = Pt(20)
    body.top = Inches(1.9); body.height = Inches(3.6)
    box(s, 0.92, 5.55, 11.5, 1.25, GREENSOFT)
    icon(s, "check", 1.15, 5.8, 0.72, GREEN)
    text(s, 2.1, 5.72, 10.1, 0.95, [[("Everything referenced here is inside the tagged commit ", {"bold": True}), (TAG, {"bold": True, "color": GREEN}), (": code, README, Dockerfile, results.jsonl, metrics.md, this deck (PPTX + PDF) and the video link.", {})]], size=14, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "Confirm each checklist item and the release tag.")

    # ---------- 12. Thank you ----------
    s = S[11]
    text(s, 1.14, 5.05, 11.2, 1.2, [[("FixFlow · Theme 02 · ", {"bold": True, "color": VIOLET}), (f"{TEAM['team']}", {"bold": True})],
                                    f"{m[0]['name']} · {m[1]['name']} · {m[2]['name']} · {m[3]['name']}",
                                    [("Repository: ", {"bold": True}), (repo_short, {}), ("   Tag: ", {"bold": True}), (TAG, {})]], size=14, color=INK)
    s.notes_slide.notes_text_frame.text = f"Presenter: {m[3]['name']}\nClose: thank the jury, repeat the one-liner, invite questions."

    prs.save(str(OUT))
    print("wrote", OUT)


if __name__ == "__main__":
    main()
